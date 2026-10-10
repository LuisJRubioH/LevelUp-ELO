"""Teacher exports open as data in a spreadsheet: text a student typed that looks like a formula
stays text, a character a workbook cannot hold does not break the download, and the CSV declares
its encoding. Both engines."""

import csv
import io
import uuid

import openpyxl

from tests.integration.conftest import (
    answer,
    enroll,
    headers_for,
    make_group,
    make_student,
    make_teacher,
    sql,
)

COURSE = "algebra_basica"


def _class_with_typed_text(repo):
    teacher = make_teacher(repo)
    sql(repo, "UPDATE users SET approved = 1 WHERE id = ?", (teacher,))
    group = make_group(repo, teacher, COURSE)
    student = make_student(repo, "colegio")
    username = f"=1+2_{uuid.uuid4().hex[:8]}"
    sql(
        repo, "UPDATE users SET username = ?, group_id = ? WHERE id = ?", (username, group, student)
    )
    enroll(repo, student, COURSE, group)
    item = repo.get_items_from_db(course_id=COURSE)[0]
    answer(repo, student, item["id"])
    repo.save_katia_interaction(
        student, COURSE, item["id"], item["topic"], "-5x + 3\x07 =SUM(A1)", "¿Qué pasa con x?"
    )
    return headers_for(repo, teacher), username


def test_the_workbook_keeps_typed_text_as_text(repo, client):
    headers, username = _class_with_typed_text(repo)

    response = client.get("/api/teacher/export/xlsx", headers=headers)

    assert response.status_code == 200, response.text
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    names = [cell for row in book["Intentos"].iter_rows() for cell in row if cell.value == username]
    assert names and all(cell.data_type == "s" for cell in names)
    katia = book["KatIA"]
    header = [cell.value for cell in katia[1]]
    column = header.index("student_message") + 1
    typed = [katia.cell(row=r, column=column).value for r in range(2, katia.max_row + 1)]
    assert "-5x + 3 =SUM(A1)" in typed


def test_the_csv_neutralises_formulas_and_declares_utf8(repo, client):
    headers, username = _class_with_typed_text(repo)

    response = client.get("/api/teacher/export/csv", headers=headers)

    assert response.status_code == 200, response.text
    body = response.content.decode("utf-8")
    assert body.startswith("﻿")
    rows = list(csv.DictReader(io.StringIO(body.lstrip("﻿"))))
    assert {row["username"] for row in rows} == {"'" + username}
