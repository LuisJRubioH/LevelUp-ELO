"""
Spec 007 delivery S (FR-070, US7-AS1): the League's answers (`pvp_answers`) in the student's and
the teacher's statistics, counted per topic and per item and marked as League answers — beside
the practice figures, never inside them. Both engines; the US7-AS1 tests play a real match on
uvicorn (see test_pvp_enrolment.py for why a match needs one event loop).
"""

import io
import json
import time

from tests.integration.conftest import (
    answer,
    enroll,
    headers_for,
    make_course,
    make_group,
    make_student,
    make_teacher,
    sql,
)
from tests.integration.test_pvp_enrolment import (  # noqa: F401 — live_api is a fixture
    ANSWER_PACE_S,
    _token,
    _until,
    live_api,
)

TOPIC = "Fracciones"


def _approved_teacher(repo) -> int:
    teacher = make_teacher(repo)
    sql(repo, "UPDATE users SET approved = 1 WHERE id = ?", (teacher,))
    return teacher


def _linked_student(repo, teacher, course_id) -> int:
    """A student enrolled in `course_id` through the teacher's group for it (spec 006 FR-040)."""
    student = make_student(repo)
    enroll(repo, student, course_id, make_group(repo, teacher, course_id))
    return student


def _closed_match(repo, course_id, a, b, answers, reason=None) -> int:
    """A finished match with `answers` = [(user_id, item_id, is_correct)], closed with `reason`
    (None: a normal finish)."""
    items = sorted({item for _, item, _ in answers})
    match_id = repo.create_pvp_match(course_id, a, b, items)
    for user_id, item_id, correct in answers:
        repo.save_pvp_answer(match_id, user_id, item_id, correct)
    repo.finish_pvp_match(
        match_id=match_id, winner_id=a, score_p1=1, score_p2=0,
        elo_delta_p1=0.0, elo_delta_p2=0.0, p1_id=a, p2_id=b, reason=reason,
    )  # fmt: skip
    return match_id


def _answered(rows) -> list:
    return sorted((r["match_id"], r["item_id"], bool(r["is_correct"])) for r in rows)


# ── Repositories: raw rows of closed matches (T011) ─────────────────────────


def test_spec007_league_answers_are_raw_rows_of_closed_matches(repo):
    """FR-070, R15: one row per answer — the student's own, never the opponent's; an active
    match's answers are not read; the course is the match's and the topic the item's."""
    course_id, items = make_course(repo, [TOPIC, "Decimales"], name="Aritmética S")
    f1, d1 = items[TOPIC][0], items["Decimales"][0]
    a, b = make_student(repo), make_student(repo)
    finished = _closed_match(repo, course_id, a, b, [(a, f1, False), (b, f1, True), (a, d1, True)])
    active = repo.create_pvp_match(course_id, a, b, [f1])
    repo.save_pvp_answer(active, a, f1, True)

    rows = repo.get_league_answers(a)

    assert _answered(rows) == sorted([(finished, f1, False), (finished, d1, True)])
    by_item = {r["item_id"]: r for r in rows}
    assert by_item[f1]["course_id"] == course_id
    assert by_item[f1]["course_name"] == "Aritmética S"
    assert by_item[f1]["topic"] == TOPIC
    assert by_item[d1]["topic"] == "Decimales"
    assert all(len(str(r["answered_at"])) == 19 for r in rows)  # 'YYYY-MM-DD HH:MM:SS'
    assert _answered(repo.get_league_answers(b)) == [(finished, f1, True)]


def test_spec007_answers_of_every_closed_match_count(repo):
    """FR-070, O-S2 (a): every closed match counts whatever its outcome — finished, forfeit,
    over the daily cap, not saved, orphaned; an active one does not."""
    course_id, items = make_course(repo, [TOPIC], items_per_topic=6)
    a, b = make_student(repo), make_student(repo)
    ids = items[TOPIC]
    finished = _closed_match(repo, course_id, a, b, [(a, ids[0], True)])
    forfeit = _closed_match(repo, course_id, a, b, [(a, ids[1], False)], reason="forfeit")
    capped = _closed_match(repo, course_id, a, b, [(a, ids[2], True)], reason="daily_cap")

    not_saved = repo.create_pvp_match(course_id, a, b, [ids[3]])
    repo.save_pvp_answer(not_saved, a, ids[3], False)
    assert repo.abandon_pvp_match(not_saved, "not_saved")

    orphaned = repo.create_pvp_match(course_id, a, b, [ids[4]])
    repo.save_pvp_answer(orphaned, a, ids[4], True)
    sql(
        repo,
        "UPDATE pvp_matches SET started_at = ? WHERE id = ?",
        ("2000-01-01 00:00:00", orphaned),
    )
    repo.expire_stale_pvp_matches()

    active = repo.create_pvp_match(course_id, a, b, [ids[5]])
    repo.save_pvp_answer(active, a, ids[5], True)

    assert _answered(repo.get_league_answers(a)) == sorted(
        [
            (finished, ids[0], True),
            (forfeit, ids[1], False),
            (capped, ids[2], True),
            (not_saved, ids[3], False),
            (orphaned, ids[4], True),
        ]
    )


def test_spec007_teacher_league_answers_follow_the_student_set(repo):
    """FR-070, spec 006 FR-040, O-S5: the teacher reads their active, non-test students' answers
    in the courses that link them — not another teacher's student, not a course that does not
    link the student, not a test student; `student_id` narrows it to one student."""
    course_id, items = make_course(repo, [TOPIC])
    other_course, other_items = make_course(repo, [TOPIC])
    item, other_item = items[TOPIC][0], other_items[TOPIC][0]
    teacher, stranger = _approved_teacher(repo), _approved_teacher(repo)
    mine = _linked_student(repo, teacher, course_id)
    enroll(repo, mine, other_course)  # enrolled, but not through this teacher
    classmate = _linked_student(repo, teacher, course_id)
    tester = _linked_student(repo, teacher, course_id)
    sql(repo, "UPDATE users SET is_test_user = 1 WHERE id = ?", (tester,))
    theirs = _linked_student(repo, stranger, course_id)

    linked = _closed_match(
        repo, course_id, mine, classmate, [(mine, item, False), (classmate, item, True)]
    )
    _closed_match(
        repo, other_course, mine, theirs, [(mine, other_item, False), (theirs, other_item, True)]
    )
    _closed_match(repo, course_id, tester, theirs, [(tester, item, False), (theirs, item, False)])

    rows = repo.get_teacher_league_answers(teacher)

    assert sorted((r["user_id"], r["match_id"], r["item_id"]) for r in rows) == sorted(
        [(mine, linked, item), (classmate, linked, item)]
    )
    names = {r["user_id"]: r["student_name"] for r in rows}
    assert names[mine] == sql(repo, "SELECT username FROM users WHERE id = ?", (mine,))[0][0]
    assert [r["user_id"] for r in repo.get_teacher_league_answers(teacher, mine)] == [mine]
    assert [r["user_id"] for r in repo.get_teacher_league_answers(stranger)] == [theirs]


# ── The export (T013) ────────────────────────────────────────────────────────


def _workbook(repo, client, teacher):
    import openpyxl

    response = client.get("/api/teacher/export/xlsx", headers=headers_for(repo, teacher))
    assert response.status_code == 200, response.text
    return openpyxl.load_workbook(io.BytesIO(response.content))


def test_spec007_the_export_has_a_league_sheet(repo, client):
    """FR-070, O-S4: the workbook gains «Liga» after «KatIA», one row per student × course ×
    topic × item with answers and correct; the four existing sheets keep their place and their
    headers."""
    course_id, items = make_course(repo, [TOPIC], name="Aritmética S")
    item = items[TOPIC][0]
    teacher = _approved_teacher(repo)
    student, rival = _linked_student(repo, teacher, course_id), make_student(repo)
    answer(repo, student, item, correct=True)  # a practice attempt: «Intentos» has a row
    _closed_match(repo, course_id, student, rival, [(student, item, False)])
    _closed_match(repo, course_id, student, rival, [(student, item, True)])

    book = _workbook(repo, client, teacher)

    assert book.sheetnames == ["Intentos", "Matrículas", "Procedimientos", "KatIA", "Liga"]
    intentos = [c.value for c in book["Intentos"][1]]
    assert intentos == list(repo.export_teacher_student_data(teacher)[0].keys())
    liga = [[c.value for c in row] for row in book["Liga"].iter_rows()]
    name = sql(repo, "SELECT username FROM users WHERE id = ?", (student,))[0][0]
    assert liga == [
        [
            "student_id",
            "name",
            "course_id",
            "course_name",
            "topic",
            "item_id",
            "answers",
            "correct",
        ],
        [student, name, course_id, "Aritmética S", TOPIC, item, 2, 1],
    ]


# ── US7-AS1: a real match, then both screens (T012) and the guard (T003) ─────


def _play(repo, live_api, course_id, first, second, miss_first=True):
    """`first` and `second` play one match in `course_id`; `first` misses the first item when
    `miss_first`, every other answer is right. Returns the items served."""
    from websockets.sync.client import connect

    url = f"{live_api}/api/ws/pvp/{course_id}"
    with connect(url) as waiting, connect(url) as joining:
        waiting.send(json.dumps({"token": _token(repo, first)}))
        _until(waiting, "waiting")
        joining.send(json.dumps({"token": _token(repo, second)}))
        items = [_until(socket, "game_start") for socket in (waiting, joining)][0]["items"]
        for socket, misses in ((waiting, miss_first), (joining, False)):
            for n, item in enumerate(items):
                time.sleep(ANSWER_PACE_S)  # spec 007 FR-051: a person's pace
                selected = "B" if misses and n == 0 else "A"
                socket.send(
                    json.dumps({"type": "answer", "item_id": item["id"], "selected": selected})
                )
                assert _until(socket, "answer_result")["item_id"] == item["id"]
        for socket in (waiting, joining):
            _until(socket, "game_end")
    return items


def _no_key_or_difficulty(text):
    assert "correct_option" not in text
    assert "difficulty" not in text


def test_spec007_a_league_miss_reaches_the_students_statistics(repo, live_api, client):
    """US7-AS1, FR-070: after a real match the student's League statistics count the miss under
    its topic and among the most-missed items, marked as League answers; the statement is
    there to study, the key and the difficulty are not."""
    course_id, _ = make_course(repo, [TOPIC], name="Aritmética S", items_per_topic=5)
    first, second = make_student(repo), make_student(repo)
    enroll(repo, first, course_id)
    enroll(repo, second, course_id)
    items = _play(repo, live_api, course_id, first, second)

    response = client.get("/api/student/league/stats", headers=headers_for(repo, first))

    assert response.status_code == 200, response.text
    body = response.json()
    assert (body["answers"], body["correct"]) == (5, 4)
    assert [
        (t["course_id"], t["topic"], t["answers"], t["correct"], t["source"])
        for t in body["topics"]
    ] == [(course_id, TOPIC, 5, 4, "league")]
    missed = body["items"]
    assert [(i["item_id"], i["answers"], i["missed"], i["source"]) for i in missed] == [
        (items[0]["id"], 1, 1, "league")
    ]
    assert missed[0]["content"] == items[0]["content"]
    assert missed[0]["options"] == ["A", "B", "C", "D"]
    _no_key_or_difficulty(response.text)
    rival = client.get("/api/student/league/stats", headers=headers_for(repo, second)).json()
    assert (rival["answers"], rival["correct"], rival["items"]) == (5, 5, [])


def test_spec007_a_league_miss_reaches_the_teachers_view(repo, live_api, client):
    """US7-AS1, FR-070: the same miss in the teacher's view of the student; another teacher gets
    404 and a student 403 (spec 006 FR-042)."""
    course_id, _ = make_course(repo, [TOPIC], name="Aritmética S", items_per_topic=5)
    teacher, stranger = _approved_teacher(repo), _approved_teacher(repo)
    first, second = _linked_student(repo, teacher, course_id), make_student(repo)
    enroll(repo, second, course_id)
    items = _play(repo, live_api, course_id, first, second)
    path = f"/api/teacher/student/{first}/league"

    response = client.get(path, headers=headers_for(repo, teacher))

    assert response.status_code == 200, response.text
    body = response.json()
    assert (body["answers"], body["correct"]) == (5, 4)
    assert [(t["topic"], t["answers"], t["correct"], t["source"]) for t in body["topics"]] == [
        (TOPIC, 5, 4, "league")
    ]
    assert [(i["item_id"], i["missed"], i["source"]) for i in body["items"]] == [
        (items[0]["id"], 1, "league")
    ]
    assert body == client.get("/api/student/league/stats", headers=headers_for(repo, first)).json()
    _no_key_or_difficulty(response.text)
    assert client.get(path, headers=headers_for(repo, stranger)).status_code == 404
    assert client.get(path, headers=headers_for(repo, second)).status_code == 403


def _figures(repo, client, teacher, student):
    """Every `attempts` figure S must leave as it is (O-S1 (a))."""
    import openpyxl

    student_headers, teacher_headers = headers_for(repo, student), headers_for(repo, teacher)
    stats = client.get("/api/student/stats", headers=student_headers).json()
    xlsx = client.get("/api/teacher/export/xlsx", headers=teacher_headers)
    book = openpyxl.load_workbook(io.BytesIO(xlsx.content))
    return {
        "total_attempts": stats["total_attempts"],
        "study_streak": stats["study_streak"],
        "subtopics": stats["subtopics"],
        "dashboard": client.get("/api/teacher/dashboard", headers=teacher_headers).json(),
        "student_panel": client.get(
            f"/api/teacher/student/{student}", headers=teacher_headers
        ).json(),
        "metrics": client.get("/api/teacher/metrics", headers=teacher_headers).json(),
        "intentos": [[c.value for c in row] for row in book["Intentos"].iter_rows()],
        "attempts_rows": sql(repo, "SELECT COUNT(*) FROM attempts")[0][0],
        "ratings": sql(
            repo,
            "SELECT user_id, course_id, topic, current_elo, rd FROM student_course_topic_elo"
            " ORDER BY user_id, course_id, topic",
        ),
    }


def test_spec007_league_answers_leave_the_practice_figures_unchanged(repo, live_api, client):
    """FR-070 «marked as League answers», O-S1 (a), R15: a League match adds no attempt and
    moves no practice figure — attempts, streak, subtopic report, the teacher's dashboard,
    student panel and metrics, the «Intentos» sheet. Neither player has a rated topic in the
    match's course, so the match applies 0 and no rating row may change at all."""
    practice, practice_items = make_course(repo, [TOPIC], name="Práctica S", items_per_topic=2)
    league, _ = make_course(repo, ["Decimales"], name="Liga S", items_per_topic=5)
    teacher = _approved_teacher(repo)
    first = _linked_student(repo, teacher, practice)
    enroll(repo, first, league, make_group(repo, teacher, league))
    second = make_student(repo)
    enroll(repo, second, league)
    answer(repo, first, practice_items[TOPIC][0], correct=True)
    answer(repo, first, practice_items[TOPIC][1], correct=False)
    before = _figures(repo, client, teacher, first)
    assert before["total_attempts"] == 2

    _play(repo, live_api, league, first, second)

    assert _figures(repo, client, teacher, first) == before
    # The match did write its answers — to pvp_answers, not to attempts.
    assert sql(repo, "SELECT COUNT(*) FROM pvp_answers WHERE user_id = ?", (first,))[0][0] == 5
