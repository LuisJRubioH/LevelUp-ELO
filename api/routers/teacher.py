"""
api/routers/teacher.py
======================
Endpoints del panel del docente:
  GET  /teacher/dashboard         → resumen de grupos y estudiantes
  GET  /teacher/groups            → grupos del docente
  POST /teacher/groups            → crear grupo
  POST /teacher/groups/{id}/invitations → invitaciones de un solo uso (spec 001 INV)
  GET  /teacher/groups/{id}/invitations → invitaciones del grupo y su estado
  DELETE /teacher/invitations/{id}      → anular una invitación sin usar
  GET  /teacher/procedures        → cola de procedimientos pendientes
  POST /teacher/procedures/grade  → calificar procedimiento
  GET  /teacher/student/{id}      → reporte detallado de un estudiante
  GET  /teacher/export            → descarga CSV/XLSX de datos
"""

import io
from datetime import datetime, timezone
from typing import Literal

from fastapi import APIRouter, Body, Depends, HTTPException, Query, Request, Response, status
from fastapi.responses import StreamingResponse

from api.ai_target import TIMED_OUT, UNAVAILABLE, ai_target
from api.dependencies import CurrentUser, RepoDep, require_role
from api.schemas.teacher import (
    ApproveTeacherRequest,
    ChangeGroupRequest,
    CreateGroupRequest,
    CreateInvitationsRequest,
    CreatedInvitationsResponse,
    ExamAssignmentCreateRequest,
    ExamAssignmentResponse,
    ExamTemplateCreateRequest,
    ExamTemplatePatchRequest,
    ExamTemplateResponse,
    GradeRequest,
    GradeResponse,
    GroupResponse,
    InvitationsResponse,
    InvitationView,
    ItemCatalogEntry,
    PendingProcedure,
    ProcedureCopy,
    ProcedureGradingEntry,
    StudentPathResponse,
    StudentReportResponse,
    StudentAIAnalysisRequest,
    UserAdminRow,
)
from src.application.services.teacher_service import TeacherService
from src.domain.exams.results import template_results
from src.domain.invitations import format_token
from src.domain.learning.prealgebra import PREALGEBRA_COURSE_ID, path_record
from src.domain.procedures import ReviewRoundSuperseded
from api.config import settings
from api.rate_limit import limiter

router = APIRouter(
    prefix="/teacher",
    tags=["teacher"],
    dependencies=[require_role("teacher")],  # spec 006 FR-042
)


def _svc(repo) -> TeacherService:
    """Composición: la capa de entrada es la que conoce infrastructure (R2)."""
    from src.infrastructure.external_api.ai_client import get_pedagogical_analysis

    return TeacherService(repository=repo, pedagogical_analysis=get_pedagogical_analysis)


def _require_teacher_student(repo, student_id: int, user: dict) -> dict:
    """The student as this teacher sees them, with `group_id` set to the teacher's group.

    Same students as the dashboard: those whose main group is the teacher's, and those enrolled
    through one of the teacher's groups (joining a second group by code replaces the main one).
    """
    student = repo.get_user_by_id(student_id)
    if not student or student.get("role") != "student":
        raise HTTPException(status_code=404, detail="Estudiante no encontrado.")
    own = {group["group_id"] for group in repo.get_groups_by_teacher(user["user_id"])}
    candidates = [student.get("group_id")] + [
        enrolment.get("group_id") for enrolment in repo.get_user_enrollments(student_id)
    ]
    group_id = next((group for group in candidates if group in own), None)
    if group_id is None:
        raise HTTPException(status_code=404, detail="Estudiante no encontrado.")
    return {**student, "group_id": group_id}


# ── Dashboard ─────────────────────────────────────────────────────────────────


Window = Literal["7", "30", "all"]


def _days(window: str) -> int | None:
    """Spec 006 FR-043: the last 7 or 30 days, or the whole history (None)."""
    return None if window == "all" else int(window)


@router.get("/dashboard")
def dashboard(user: CurrentUser, repo: RepoDep, days: Window = "30"):
    """The teacher's students (spec 006 FR-040) with their attempts in the window (FR-043), their
    current ratings, the groups and the server's summary (FR-041)."""
    svc = _svc(repo)
    students = repo.get_teacher_dashboard_stats(user["user_id"], days=_days(days))
    views = svc.ratings.ratings_view_bulk([s["user_id"] for s in students])
    for s in students:  # the overall rating as the student sees it (FR-028a, FR-028j)
        view = views[s["user_id"]]
        s.update(
            global_elo=view["overall"],
            display_rating=view["display_rating"],
            rank_label=view["rank_label"],
            overall_status=view["overall_status"],
        )
    memberships = repo.get_teacher_student_groups(user["user_id"])
    summary = svc.ratings.console_summary(
        [
            {
                "user_id": s["user_id"],
                "accuracy": s["accuracy"],
                "total_attempts": s["total_attempts"],
                "group_ids": memberships.get(s["user_id"], []),
            }
            for s in students
        ]
    )
    groups = repo.get_groups_by_teacher(user["user_id"])
    # The group's reusable code (V1 still reads it) is not sent: spec 001 FR-028u.
    groups = [{k: v for k, v in g.items() if k != "invite_code"} for g in groups]
    return {"students": students, "groups": groups, "summary": summary}


# ── Grupos ────────────────────────────────────────────────────────────────────


@router.get("/groups", response_model=list[GroupResponse])
def get_groups(user: CurrentUser, repo: RepoDep):
    """Lista de grupos del docente con contador de estudiantes."""
    rows = repo.get_groups_by_teacher(user["user_id"])
    result = []
    for r in rows:
        row = (
            dict(r)
            if isinstance(r, dict)
            else {
                "id": r[0],
                "name": r[1],
                "course_id": r[2] if len(r) > 2 else None,
                "student_count": r[4] if len(r) > 4 else 0,
            }
        )
        result.append(
            GroupResponse(
                group_id=row.get("group_id") or row.get("id"),  # both keys present
                name=row["name"],
                course_id=row.get("course_id"),
                student_count=row.get("student_count", 0),
            )
        )
    return result


@router.post("/groups", status_code=status.HTTP_201_CREATED, response_model=GroupResponse)
def create_group(body: CreateGroupRequest, user: CurrentUser, repo: RepoDep):
    """Crea un nuevo grupo para el docente."""
    svc = _svc(repo)
    ok, msg, group_id = svc.create_new_group(user["user_id"], body.course_id, body.group_name)
    if not ok:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=msg)
    return GroupResponse(
        group_id=group_id,
        name=body.group_name,
        course_id=body.course_id,
        student_count=0,
    )


# ── Invitaciones de un solo uso (spec 001, delivery INV) ──────────────────────
# The repository answers None / "not_found" for a group or an invitation that is not this
# teacher's (an admin has none): 404, revealing nothing about it (FR-028p, FR-028t).


@router.post(
    "/groups/{group_id}/invitations",
    status_code=status.HTTP_201_CREATED,
    response_model=CreatedInvitationsResponse,
)
def create_invitations(
    group_id: int, body: CreateInvitationsRequest, user: CurrentUser, repo: RepoDep
):
    """FR-028p: 1–50 single-use invitations to the teacher's own group."""
    ttl_days = settings.invitation_ttl_days
    tokens = repo.create_group_invitations(user["user_id"], group_id, body.count, ttl_days)
    if tokens is None:
        raise HTTPException(status_code=404, detail="Grupo no encontrado.")
    return CreatedInvitationsResponse(tokens=[format_token(t) for t in tokens], ttl_days=ttl_days)


@router.get("/groups/{group_id}/invitations", response_model=InvitationsResponse)
def list_invitations(group_id: int, user: CurrentUser, repo: RepoDep):
    """FR-028t: the group's invitations, newest first, with their state."""
    rows = repo.get_group_invitations(user["user_id"], group_id)
    if rows is None:
        raise HTTPException(status_code=404, detail="Grupo no encontrado.")
    return InvitationsResponse(
        ttl_days=settings.invitation_ttl_days,
        invitations=[InvitationView(**{**r, "token": format_token(r["token"])}) for r in rows],
    )


@router.delete("/invitations/{invitation_id}", status_code=status.HTTP_204_NO_CONTENT)
def revoke_invitation(invitation_id: int, user: CurrentUser, repo: RepoDep):
    """FR-028t: revoke an unused, unexpired invitation of the teacher's own group."""
    outcome = repo.revoke_group_invitation(user["user_id"], invitation_id)
    if outcome == "not_found":
        raise HTTPException(status_code=404, detail="Invitación no encontrada.")
    if outcome != "revoked":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Solo se puede anular una invitación sin usar y vigente.",
        )


# ── Procedimientos ────────────────────────────────────────────────────────────


@router.get("/procedures", response_model=list[PendingProcedure])
def pending_procedures(user: CurrentUser, repo: RepoDep):
    """Cola de procedimientos pendientes de calificación."""
    rows = repo.get_pending_submissions_for_teacher(user["user_id"])
    result = []
    for r in rows:
        row = dict(r) if isinstance(r, dict) else {}
        if not row:
            continue
        has_image = bool(
            row.get("storage_url") or row.get("image_data") or row.get("procedure_image_path")
        )
        result.append(
            PendingProcedure(
                submission_id=row.get("id", 0),
                student_id=row.get("student_id") or row.get("user_id", 0),
                student_name=row.get("student_name") or "",
                item_id=row.get("item_id", ""),
                item_content=row.get("item_content"),
                ai_score=row.get("ai_proposed_score"),
                ai_feedback=row.get("ai_feedback"),
                status=row.get("status", "pending"),
                created_at=str(row.get("submitted_at") or row.get("created_at", "")),
                has_image=has_image,
                review_round=row.get("review_round") or 1,
            )
        )
    return result


@router.get("/procedures/copies", response_model=list[ProcedureCopy])
def procedure_copies(user: CurrentUser, repo: RepoDep):
    """Spec 001 FR-024i: procedures graded 0 automatically as copies in this teacher's groups,
    last 30 days. Never the student who sent the file first."""
    return [
        ProcedureCopy(
            submission_id=r["submission_id"],
            student_id=r["student_id"],
            student_name=r["student_name"] or "",
            item_id=r["item_id"],
            item_content=r.get("item_content"),
            reviewed_at=str(r["reviewed_at"]) if r.get("reviewed_at") else None,
        )
        for r in repo.get_recent_procedure_copies_for_teacher(user["user_id"])
    ]


@router.get("/procedures/{submission_id}/gradings", response_model=list[ProcedureGradingEntry])
def procedure_gradings(submission_id: int, user: CurrentUser, repo: RepoDep):
    """Spec 001 FR-024e: every grading of a procedure, for the teacher of the student's group."""
    rows = repo.get_procedure_gradings(submission_id, teacher_id=user["user_id"])
    if rows is None:
        raise HTTPException(status_code=404, detail="Procedimiento no encontrado.")
    fields = ProcedureGradingEntry.model_fields
    return [ProcedureGradingEntry(**{k: r[k] for k in fields}) for r in rows]


@router.get("/procedures/{submission_id}/image")
def procedure_image(submission_id: int, user: CurrentUser, repo: RepoDep):
    """Retorna la imagen de un procedimiento (desde Supabase Storage o BYTEA)."""
    from fastapi.responses import Response as FastAPIResponse

    rows = repo.get_pending_submissions_for_teacher(user["user_id"])
    row = next(
        (
            dict(r)
            for r in rows
            if (dict(r) if isinstance(r, dict) else {}).get("id") == submission_id
        ),
        None,
    )
    if row is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Procedimiento no encontrado."
        )

    image_bytes: bytes | None = None
    mime = row.get("mime_type") or "image/jpeg"

    # Intentar desde Supabase Storage primero
    storage_url = row.get("storage_url")
    if storage_url and hasattr(repo, "resolve_storage_image"):
        try:
            image_bytes = repo.resolve_storage_image(storage_url)
        except Exception:
            pass

    # Fallback a BYTEA en DB. The queue leaves image_data out when the row has a storage_url
    # (so it does not load every image); read it for this one submission if Storage failed.
    if not image_bytes:
        image_bytes = row.get("image_data") or repo.get_submission_image_data(submission_id)

    if not image_bytes:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Imagen no disponible.")

    return FastAPIResponse(content=bytes(image_bytes), media_type=mime)


@router.post("/procedures/grade", response_model=GradeResponse)
def grade_procedure(body: GradeRequest, user: CurrentUser, repo: RepoDep):
    """Califica un procedimiento y aplica al rating solo la diferencia con su contribución
    anterior (spec 001 bloque 7, FR-024d). Un reintento con la misma `review_round` devuelve el
    mismo resultado sin aplicar nada (FR-024g); una ronda que el estudiante reemplazó al reenviar
    responde 409 (FR-024l)."""
    # Usar la misma cola autorizada que el panel/visor; no confiar en el ID enviado por el
    # cliente. Fuera de la cola solo se admite el reintento de una ronda ya calificada, que el
    # repositorio comprueba contra el grupo del docente.
    submission = next(
        (
            row
            for row in repo.get_pending_submissions_for_teacher(user["user_id"])
            if row["id"] == body.submission_id
        ),
        None,
    )
    if submission is None and body.review_round is None:
        raise HTTPException(status_code=404, detail="Procedimiento pendiente no encontrado.")

    svc = _svc(repo)
    try:
        result = svc.grade_procedure(
            submission_id=body.submission_id,
            teacher_score=body.teacher_score,
            feedback=body.teacher_feedback or "",
            teacher_id=user["user_id"],
            review_round=body.review_round,
        )
    except ReviewRoundSuperseded as exc:
        # FR-024l: the student resubmitted while this review was open; nothing was graded.
        raise HTTPException(
            status_code=409,
            detail="El estudiante reenvió este procedimiento; revisa el envío nuevo antes de calificar.",
        ) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc

    # Notificar al estudiante en tiempo real (WebSocket), una sola vez por calificación.
    try:
        if submission and result["created"]:
            from api.websocket.notifications import notify_sync

            notify_sync(
                room=f"student_{submission['student_id']}",
                event="procedure_graded",
                data={
                    "submission_id": body.submission_id,
                    "teacher_score": result["grade"],
                    "elo_delta": result["applied"],
                    "computed_delta": result["computed_delta"],
                    "feedback": body.teacher_feedback or "",
                },
            )
    except Exception:
        pass  # No fallar el endpoint si la notificación falla

    return GradeResponse(
        submission_id=body.submission_id,
        teacher_score=result["grade"],
        elo_delta=result["applied"],
        status="graded",
        computed_delta=result["computed_delta"],
        contribution=result["contribution_after"],
        review_round=result["review_round"],
    )


# ── Reporte por estudiante ────────────────────────────────────────────────────


@router.get("/student/{student_id}")
def student_report(student_id: int, user: CurrentUser, repo: RepoDep):
    """Reporte detallado de un estudiante (ELO, intentos, procedimientos)."""
    _require_teacher_student(repo, student_id, user)
    svc = _svc(repo)
    return svc.get_student_dashboard(student_id)


@router.get("/student/{student_id}/path/{course_id}", response_model=StudentPathResponse)
def student_path(student_id: int, course_id: str, user: CurrentUser, repo: RepoDep):
    """Spec 003 L4 (FR-042): the student's learning path — state, every answer in order and the
    post-diagnostic outcomes. Right or wrong and the tag; never the chosen value (N11)."""
    _require_teacher_student(repo, student_id, user)
    if course_id != PREALGEBRA_COURSE_ID:
        raise HTTPException(status_code=404, detail="Este curso no tiene ruta de aprendizaje")
    return {
        "student_id": student_id,
        "course_id": course_id,
        "nodes": path_record(
            repo.get_lesson_states(student_id, course_id),
            repo.get_lesson_answers(student_id, course_id),
            repo.get_lesson_outcomes(student_id, course_id),
        ),
    }


@router.get("/student/{student_id}/elo-history")
def student_elo_history(
    student_id: int, user: CurrentUser, repo: RepoDep, limit: int = Query(20, ge=1, le=200)
):
    """Historial de ELO del estudiante para el gráfico temporal."""
    _require_teacher_student(repo, student_id, user)
    attempts = repo.get_latest_attempts(student_id, limit=limit)
    return {"attempts": list(reversed(attempts))}


@router.get("/student/{student_id}/katia-history")
def student_katia_history(student_id: int, user: CurrentUser, repo: RepoDep):
    """Spec 005 FR-033: the student's closed KatIA conversations — topic, item, turns, outcome,
    kind of difficulty and tag; never what was written."""
    _require_teacher_student(repo, student_id, user)
    return {"conversations": repo.get_katia_conversations(student_id)}


@router.post("/student/{student_id}/ai-analysis")
@limiter.limit(settings.rate_limit_socratic)
def student_ai_analysis(
    student_id: int,
    request: Request,
    user: CurrentUser,
    repo: RepoDep,
    body: StudentAIAnalysisRequest = Body(default_factory=StudentAIAnalysisRequest),
):
    """Genera un análisis pedagógico del estudiante. Key: docente > AI_KEY_TEACHER_ANALYSIS > general."""
    _require_teacher_student(repo, student_id, user)
    from api.config import settings

    svc = _svc(repo)
    effective_key = settings.get_ai_key("teacher_analysis", body.api_key)
    # Igual que KatIA y la revisión de procedimientos: sin key no hay proveedor. Antes
    # caía al servidor local (localhost:1234), su ConnectionError llegaba como 500 sin
    # cabeceras CORS y el navegador lo mostraba como un error de red.
    if not effective_key:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="No hay API key de IA configurada. Pide al administrador que configure "
            "SYSTEM_AI_API_KEY.",
        )
    # Spec 005 FR-010 … FR-013: the key's prefix picks the provider, settings the model.
    target = ai_target("teacher_analysis", effective_key)
    global_elo = svc.ratings.ratings_view(student_id)["overall"]
    # FR-015: the student's graded procedures (0-100), overall and per course.
    scores = [row["score"] for row in repo.get_student_procedure_scores(student_id)]
    procedure_stats = (
        {"count": len(scores), "avg_score": sum(scores) / len(scores), "scores": scores}
        if scores
        else None
    )

    try:
        analysis = svc.generate_ai_analysis(
            student_id=student_id,
            global_elo=global_elo,
            api_key=target.api_key,
            provider=target.provider,
            base_url=target.base_url,
            model_name=target.model,
            timeout=target.timeout,
            procedure_stats=procedure_stats,
            procedure_stats_by_course=repo.get_procedure_stats_by_course(student_id),
        )
    except TimeoutError:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail=TIMED_OUT)
    except ConnectionError:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail=UNAVAILABLE)
    # FR-016: a provider's error text («ERROR_401: …», «❌ …») is not an analysis.
    if not (analysis or "").strip() or analysis.lstrip().startswith(("ERROR_", "❌")):
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail="La IA no pudo generar el análisis. Intenta de nuevo.",
        )
    return {"analysis": analysis}


@router.get("/metrics")
def teacher_metrics(user: CurrentUser, repo: RepoDep, days: Window = "30"):
    """Time per question, abandonment, top topics and activity of the teacher's students, in the
    window (spec 006 FR-040, FR-043)."""
    return repo.get_teacher_metrics(user["user_id"], days=_days(days))


@router.get("/student/{student_id}/ranking")
def student_group_ranking(student_id: int, user: CurrentUser, repo: RepoDep):
    """Ranking del grupo al que pertenece el estudiante."""
    from api.routers.student import ranking_response

    student = _require_teacher_student(repo, student_id, user)
    group_id = student.get("group_id")
    if not group_id:
        return {"basis": None, "ranking": [], "my_rank": None}
    # No requested course here: the group's course, else the overall rating (FR-028d).
    view = _svc(repo).ratings.ranking_view("group", group_id=group_id)
    return ranking_response(repo, view, student_id)


# ── Exportación ───────────────────────────────────────────────────────────────

# Exported cells hold text students typed (usernames, KatIA messages). A spreadsheet runs a cell
# that starts like a formula, so text starting with one of these is written as text (OWASP CSV
# injection); in the workbook only "=" makes a formula.
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def _csv_cell(value):
    """Text that would start a formula gets a leading apostrophe, which Excel shows as text."""
    if isinstance(value, str) and value.startswith(_FORMULA_STARTS):
        return "'" + value
    return value


def _append_xlsx_row(ws, values) -> None:
    """Append a row as data: characters a workbook cannot hold are dropped, and text starting
    with "=" is stored as text instead of a formula."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    ws.append(
        [ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in values]
    )
    for cell in ws[ws.max_row]:
        if cell.data_type == "f":
            cell.data_type = "s"


@router.get("/export/csv")
def export_csv(user: CurrentUser, repo: RepoDep):
    """Exporta los datos de intentos de los estudiantes del docente como CSV."""
    import csv
    import io

    rows = repo.export_teacher_student_data(user["user_id"])
    if not rows:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Sin datos para exportar."
        )

    output = io.StringIO()
    # The byte-order mark tells Excel the file is UTF-8, so accents read correctly.
    output.write("\ufeff")
    keys = list(rows[0].keys()) if isinstance(rows[0], dict) else []
    writer = csv.DictWriter(output, fieldnames=keys)
    writer.writeheader()
    writer.writerows([{k: _csv_cell(v) for k, v in dict(r).items()} for r in rows])

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=levelup_data.csv"},
    )


@router.get("/export/xlsx")
def export_xlsx(user: CurrentUser, repo: RepoDep):
    """Exporta los datos de los estudiantes como Excel con 4 hojas."""
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="openpyxl no instalado en el servidor.",
        )

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    datasets = {
        "Intentos": repo.export_teacher_student_data(user["user_id"]),
        "Matrículas": repo.export_teacher_enrollments(user["user_id"]),
        "Procedimientos": repo.export_teacher_procedures(user["user_id"]),
        # Spec 005 FR-033: conversations counted per course, topic and outcome; no text.
        "KatIA": repo.export_teacher_katia_summary(user["user_id"]),
    }

    for sheet_name, rows in datasets.items():
        ws = wb.create_sheet(title=sheet_name)
        if rows:
            headers = list(rows[0].keys()) if isinstance(rows[0], dict) else []
            ws.append(headers)
            for row in rows:
                _append_xlsx_row(ws, list(dict(row).values()))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=levelup_data.xlsx"},
    )


# ── Sprint C: Plantillas de examen del docente ───────────────────────────────


_LOCKED = (
    "Ya hay estudiantes que empezaron este examen: sus preguntas no cambian. "
    "Duplica la plantilla para cambiarlas."
)


def _template_to_response(t: dict, repo=None) -> ExamTemplateResponse:
    """With `repo`, also whether the template is assigned (else a draft, spec 006 FR-030) and
    whether its items are locked (FR-031)."""
    return ExamTemplateResponse(
        id=t["id"],
        teacher_id=t["teacher_id"],
        course_id=t["course_id"],
        title=t["title"],
        time_limit_min=t["time_limit_min"],
        item_ids=t["item_ids"],
        archived=t["archived"],
        created_at=t["created_at"],
        assigned=bool(repo and repo.list_assignments_for_template(t["id"])),
        locked=bool(repo and repo.template_has_sessions(t["id"])),
    )


@router.get("/exam-templates", response_model=list[ExamTemplateResponse])
def list_exam_templates(
    user: CurrentUser,
    repo: RepoDep,
    course_id: str | None = None,
    include_archived: bool = False,
):
    """Lista las plantillas creadas por el docente actual (filtrable por curso)."""
    templates = repo.list_exam_templates(
        course_id=course_id,
        teacher_id=user["user_id"],
        include_archived=include_archived,
    )
    return [_template_to_response(t, repo) for t in templates]


@router.post(
    "/exam-templates",
    response_model=ExamTemplateResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_exam_template(body: ExamTemplateCreateRequest, user: CurrentUser, repo: RepoDep):
    """Crea una plantilla de examen manual."""
    if len(body.item_ids) != len(set(body.item_ids)):
        raise HTTPException(status_code=400, detail="La plantilla contiene items duplicados.")
    # Validar que todos los item_ids existan en el curso indicado
    bank_items = repo.get_items_from_db(course_id=body.course_id)
    valid_ids = {it["id"] for it in bank_items}
    invalid = [i for i in body.item_ids if i not in valid_ids]
    if invalid:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Items no encontrados en el curso: {invalid[:5]}",
        )
    template_id = repo.create_exam_template(
        teacher_id=user["user_id"],
        course_id=body.course_id,
        title=body.title,
        time_limit_min=body.time_limit_min,
        item_ids=body.item_ids,
    )
    created = repo.get_exam_template(template_id)
    return _template_to_response(created)


@router.patch("/exam-templates/{template_id}", response_model=ExamTemplateResponse)
def update_exam_template(
    template_id: int,
    body: ExamTemplatePatchRequest,
    user: CurrentUser,
    repo: RepoDep,
):
    """Actualiza campos puntuales de una plantilla propia."""
    template = repo.get_exam_template(template_id)
    if not template:
        raise HTTPException(status_code=404, detail="Plantilla no encontrada.")
    if template["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="No es tu plantilla.")
    if body.item_ids is not None:
        if body.item_ids != template["item_ids"] and repo.template_has_sessions(template_id):
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=_LOCKED)
        if len(body.item_ids) != len(set(body.item_ids)):
            raise HTTPException(status_code=400, detail="La plantilla contiene items duplicados.")
        bank_items = repo.get_items_from_db(course_id=template["course_id"])
        valid_ids = {it["id"] for it in bank_items}
        invalid = [i for i in body.item_ids if i not in valid_ids]
        if invalid:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Items no encontrados en el curso: {invalid[:5]}",
            )
    repo.update_exam_template(
        template_id,
        title=body.title,
        time_limit_min=body.time_limit_min,
        item_ids=body.item_ids,
    )
    return _template_to_response(repo.get_exam_template(template_id), repo)


@router.post(
    "/exam-templates/{template_id}/duplicate",
    response_model=ExamTemplateResponse,
    status_code=status.HTTP_201_CREATED,
)
def duplicate_exam_template(template_id: int, user: CurrentUser, repo: RepoDep):
    """Spec 006 FR-031: a copy, unassigned and unlocked, to change a template already taken."""
    template = _ensure_owns_template(repo, template_id, user)
    copy_id = repo.create_exam_template(
        teacher_id=user["user_id"],
        course_id=template["course_id"],
        title=f"{template['title']} (copia)",
        time_limit_min=template["time_limit_min"],
        item_ids=template["item_ids"],
    )
    return _template_to_response(repo.get_exam_template(copy_id), repo)


@router.delete("/exam-templates/{template_id}", status_code=status.HTTP_204_NO_CONTENT)
def archive_exam_template(template_id: int, user: CurrentUser, repo: RepoDep):
    """Archiva (soft delete) una plantilla propia."""
    template = repo.get_exam_template(template_id)
    if not template:
        raise HTTPException(status_code=404, detail="Plantilla no encontrada.")
    if template["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="No es tu plantilla.")
    repo.archive_exam_template(template_id)
    return Response(status_code=status.HTTP_204_NO_CONTENT)


# ── Asignación de plantillas a grupos + ventana de tiempo ────────────────────


def _as_utc(value: datetime) -> datetime:
    """A time with no zone is UTC."""
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def _stored_utc(value: datetime | None) -> str | None:
    """A window bound as naive UTC text: what both engines compare with their own clock
    (`datetime('now')` on SQLite, `NOW()` on PostgreSQL)."""
    return None if value is None else _as_utc(value).strftime("%Y-%m-%d %H:%M:%S")


def _utc_iso(value) -> str | None:
    """A stored time (naive UTC, or ISO text from older rows) as ISO 8601 with its zone, so a
    browser reads it as UTC and shows it in the teacher's time."""
    if value in (None, ""):
        return None
    if not isinstance(value, datetime):
        try:
            value = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except ValueError:
            return str(value)
    return _as_utc(value).strftime("%Y-%m-%dT%H:%M:%SZ")


def _assignment_response(row: dict) -> ExamAssignmentResponse:
    return ExamAssignmentResponse(
        **{
            **row,
            "starts_at": _utc_iso(row.get("starts_at")),
            "ends_at": _utc_iso(row.get("ends_at")),
            "created_at": _utc_iso(row.get("created_at")) or "",
        }
    )


def _ensure_owns_template(repo, template_id: int, user: dict) -> dict:
    """Carga la plantilla y verifica autoría. Lanza HTTPException."""
    template = repo.get_exam_template(template_id)
    if not template:
        raise HTTPException(status_code=404, detail="Plantilla no encontrada.")
    if template["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="No es tu plantilla.")
    return template


@router.get(
    "/exam-templates/{template_id}/assignments",
    response_model=list[ExamAssignmentResponse],
)
def list_template_assignments(template_id: int, user: CurrentUser, repo: RepoDep):
    """Lista a qué grupos está asignada esta plantilla y sus ventanas de tiempo."""
    _ensure_owns_template(repo, template_id, user)
    rows = repo.list_assignments_for_template(template_id)
    return [_assignment_response(r) for r in rows]


@router.post(
    "/exam-templates/{template_id}/assignments",
    response_model=list[ExamAssignmentResponse],
    status_code=status.HTTP_201_CREATED,
)
def create_template_assignments(
    template_id: int,
    body: ExamAssignmentCreateRequest,
    user: CurrentUser,
    repo: RepoDep,
):
    """Asigna la plantilla a uno o varios grupos. Upsert por (template, group).

    Spec 006 FR-032: the groups are the template author's and of the template's course; all of
    them are assigned in one transaction, or none."""
    template = _ensure_owns_template(repo, template_id, user)
    group_courses = {
        g["group_id"]: g["course_id"] for g in repo.get_groups_by_teacher(template["teacher_id"])
    }
    invalid = [gid for gid in body.group_ids if gid not in group_courses]
    if invalid:
        raise HTTPException(
            status_code=400,
            detail=f"Grupos no autorizados: {invalid}",
        )
    other_course = [gid for gid in body.group_ids if group_courses[gid] != template["course_id"]]
    if other_course:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Un examen se asigna solo a grupos de su curso; no se asignó a ninguno.",
        )
    repo.assign_template_to_groups(
        template_id,
        list(body.group_ids),
        starts_at=_stored_utc(body.starts_at),
        ends_at=_stored_utc(body.ends_at),
    )
    rows = repo.list_assignments_for_template(template_id)
    return [_assignment_response(r) for r in rows]


@router.delete(
    "/exam-templates/{template_id}/assignments/{assignment_id}",
    status_code=status.HTTP_204_NO_CONTENT,
)
def delete_template_assignment(
    template_id: int,
    assignment_id: int,
    user: CurrentUser,
    repo: RepoDep,
):
    """Elimina una asignación específica."""
    _ensure_owns_template(repo, template_id, user)
    if not repo.delete_exam_assignment(assignment_id, template_id=template_id):
        raise HTTPException(status_code=404, detail="Asignación no encontrada.")
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.get("/exam-templates/{template_id}/results")
def exam_template_results(template_id: int, user: CurrentUser, repo: RepoDep):
    """Análisis de resultados del examen: por pregunta, por tópico, mejor/peor
    pregunta y tema a reforzar. Los exámenes se califican automáticamente; este
    endpoint solo agrega las respuestas ya guardadas (no afecta ELO)."""
    template = _ensure_owns_template(repo, template_id, user)
    # Spec 006 FR-033: the template author's students, each on their first submission.
    students = {
        s["user_id"]: s["name"] for s in repo.get_teacher_dashboard_stats(template["teacher_id"])
    }
    analysis = template_results(repo.get_exam_template_responses(template_id), students)
    analysis["title"] = template["title"]
    return analysis


@router.get("/courses")
def list_all_courses(repo: RepoDep, user: CurrentUser):
    """Lista todos los cursos del banco — para que el docente arme exámenes."""
    courses = repo.get_courses()
    return [
        {
            "id": c["id"] if isinstance(c, dict) else c[0],
            "name": c["name"] if isinstance(c, dict) else c[1],
            "block": c["block"] if isinstance(c, dict) else c[2],
        }
        for c in courses
    ]


@router.get("/items", response_model=list[ItemCatalogEntry])
def list_items_catalog(
    repo: RepoDep,
    user: CurrentUser,
    course_id: str,
):
    """Catálogo de items de un curso — para que el docente arme exámenes."""
    items = repo.get_items_from_db(course_id=course_id)
    items.sort(key=lambda it: it["difficulty"])
    return [
        ItemCatalogEntry(
            id=it["id"],
            content=it["content"],
            difficulty=it["difficulty"],
            topic=it["topic"],
            tags=it.get("tags") or [],
        )
        for it in items
    ]
